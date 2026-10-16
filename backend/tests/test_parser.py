import pytest

from app.core.limits import ValidationError
from app.services.parser import (
    chunk_text,
    normalize_text,
    parse_document,
    parse_txt,
)


def test_normalize_text():
    raw = "  привет   мир  \n\n\n\n  как   дела  "
    result = normalize_text(raw)
    assert result == "привет мир\n\nкак дела"


def test_chunk_text_short():
    text = "короткий текст"
    chunks = chunk_text(text, max_chars=100)
    assert chunks == [text]


def test_chunk_text_long():
    text = "\n\n".join([f"Абзац номер {i} " * 20 for i in range(10)])
    chunks = chunk_text(text, max_chars=500)
    assert len(chunks) > 1
    for chunk in chunks:
        assert len(chunk) <= 600  # небольшой запас


def test_parse_txt(tmp_path):
    file = tmp_path / "test.txt"
    file.write_text("Привет, мир!", encoding="utf-8")
    assert parse_txt(file) == "Привет, мир!"


def test_parse_txt_cp1251(tmp_path):
    file = tmp_path / "test.txt"
    file.write_bytes("Привет, мир!".encode("cp1251"))
    assert parse_txt(file) == "Привет, мир!"


def test_parse_docx(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Первый абзац")
    doc.add_paragraph("Второй абзац")
    file = tmp_path / "test.docx"
    doc.save(str(file))

    result = parse_document(file)
    assert "Первый абзац" in result
    assert "Второй абзац" in result


def test_parse_pdf_invalid(tmp_path):
    file = tmp_path / "test.pdf"
    file.write_bytes(b"not a real pdf")
    with pytest.raises(ValidationError):
        parse_document(file)


def test_parse_unsupported():
    with pytest.raises(ValidationError):
        parse_document(__import__("pathlib").Path("test.xyz"))
